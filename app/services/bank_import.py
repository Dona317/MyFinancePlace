"""
Bank statement import (Excel / CSV) → transactions.

Supported layouts:
  • Fineco         — "Data_Operazione | Data_Valuta | Entrate | Uscite | Descrizione | Descrizione_Completa | Stato | Moneymap"
  • Intesa Sanpaolo — "Data | Operazione | Dettagli | Conto o carta | Contabilizzazione | Categoria | Valuta | Importo"
                      (older exports: "Data contabile | Data valuta | Descrizione | Accrediti | Addebiti | Descrizione estesa")
  • Generic        — any statement whose header has a date, a description and either a signed amount
                      or separate credit/debit columns (UniCredit, BPER, Revolut, N26, …)

Pipeline:  read_table() → detect_layout() → parse_rows() → categorize() → fingerprint()
Bank export preamble rows (account holder, period, balances) and footer rows are skipped automatically.
"""
import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser

from app.models.transaction import Transaction
from app.services.transfer import parse_amount, parse_date

HEADER_SCAN_ROWS = 40
PENDING_STATUSES = ("non contabilizzat", "autorizzat", "in attesa", "pending", "da contabilizzare")


# ── Bank layouts ───────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class BankLayout:
    key: str
    name: str
    markers: tuple[str, ...]              # text in the preamble / filename identifying the bank
    columns: dict[str, tuple[str, ...]]   # field → accepted (normalized) header names, by priority


BANKS = {
    "fineco": BankLayout(
        key="fineco",
        name="Fineco",
        markers=("fineco",),
        columns={
            "date":        ("data operazione", "data"),
            "description": ("descrizione completa", "descrizione"),
            "details":     ("descrizione",),
            "credit":      ("entrate",),
            "debit":       ("uscite",),
            "category":    ("moneymap",),
            "status":      ("stato",),
        },
    ),
    "intesa": BankLayout(
        key="intesa",
        name="Intesa Sanpaolo",
        markers=("intesa", "sanpaolo", "isybank"),
        columns={
            "date":        ("data contabile", "data operazione", "data"),
            "description": ("operazione", "descrizione"),
            "details":     ("dettagli", "descrizione estesa"),
            "amount":      ("importo",),
            "credit":      ("accrediti",),
            "debit":       ("addebiti",),
            "category":    ("categoria",),
            "currency":    ("valuta",),
            "status":      ("contabilizzazione",),
        },
    ),
    "generic": BankLayout(
        key="generic",
        name="Altra banca (generico)",
        markers=(),
        columns={
            "date": (
                "data operazione", "data contabile", "data registrazione", "data movimento", "data",
                "booking date", "transaction date", "started date", "completed date", "date",
            ),
            "description": (
                "descrizione operazione", "descrizione", "causale", "operazione", "dettaglio",
                "description", "details", "payee", "beneficiario",
            ),
            "details":  ("descrizione estesa", "descrizione completa", "dettagli", "note", "reference"),
            "amount":   ("importo eur", "importo euro", "importo", "ammontare", "amount eur", "amount"),
            "credit":   ("entrate", "accrediti", "avere", "credit", "money in"),
            "debit":    ("uscite", "addebiti", "dare", "debit", "money out"),
            "category": ("categoria", "category"),
            "currency": ("divisa", "currency"),
            "status":   ("stato", "state", "status"),
        },
    ),
}

AUTO = "auto"
DETECTION_ORDER = ("fineco", "intesa", "generic")


# ── Auto-categorization ────────────────────────────────────────────────────────

# First match wins; keywords are matched as whole words, case-insensitive.
CATEGORY_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("Stipendio",    ("stipendio", "emolumenti", "retribuzione", "salario", "busta paga", "cedolino")),
    ("Alimentari",   ("esselunga", "coop", "conad", "carrefour", "lidl", "eurospin", "pam", "penny", "aldi",
                      "supermercato", "supermercati", "iper", "despar", "famila", "bennet", "naturasi",
                      "alimentari", "spesa")),
    ("Trasporto",    ("eni", "enilive", "q8", "tamoil", "esso", "shell", "carburante", "carburanti", "benzina",
                      "trenitalia", "italo", "atm", "atac", "autostrade", "telepass", "uber", "taxi", "freenow",
                      "parcheggio", "trasporti")),
    ("Abbonamenti",  ("netflix", "spotify", "prime video", "amazon prime", "disney", "dazn", "now tv", "icloud",
                      "apple com bill", "youtube", "tim", "vodafone", "windtre", "iliad", "fastweb", "ho mobile",
                      "abbonamento", "abbonamenti")),
    ("Casa",         ("affitto", "condominio", "enel", "a2a", "hera", "iren", "edison", "sorgenia", "bolletta",
                      "tari", "ikea", "leroy merlin", "casa", "utenze")),
    ("Salute",       ("farmacia", "ospedale", "medico", "dentista", "asl", "ticket", "sanitaria", "salute")),
    ("Svago",        ("ristorante", "pizzeria", "trattoria", "bar", "cinema", "just eat", "deliveroo", "glovo",
                      "mcdonald", "burger king", "booking com", "airbnb", "svago", "tempo libero", "viaggi")),
    ("Investimenti", ("compravendita titoli", "acquisto titoli", "etf", "fondi", "pac", "dossier titoli",
                      "investimenti")),
    ("Commissioni",  ("commissioni", "commissione", "canone", "imposta di bollo", "spese tenuta conto",
                      "competenze")),
    ("Rimborsi",     ("rimborso", "storno")),
]

TRANSFER_KEYWORDS = ("giroconto", "trasferimento tra conti", "ricarica carta", "girofondi")

_RULE_PATTERNS = [
    (category, re.compile(r"\b(" + "|".join(re.escape(k) for k in keywords) + r")\b"))
    for category, keywords in CATEGORY_RULES
]
_TRANSFER_PATTERN = re.compile(r"\b(" + "|".join(re.escape(k) for k in TRANSFER_KEYWORDS) + r")\b")


def _search_text(*parts: str | None) -> str:
    return " " + re.sub(r"[^a-z0-9]+", " ", " ".join(p for p in parts if p).lower()) + " "


def categorize(description: str, details: str | None = None, bank_category: str | None = None) -> str:
    """App category for a statement row: keyword rules first, then the bank's own category, else "Altro"."""
    text = _search_text(description, details)
    for category, pattern in _RULE_PATTERNS:
        if pattern.search(text):
            return category
    if bank_category:
        bank_text = _search_text(bank_category)
        for category, pattern in _RULE_PATTERNS:
            if pattern.search(bank_text):
                return category
        return bank_category.strip()[:100]
    return "Altro"


def is_transfer(description: str, details: str | None = None) -> bool:
    return bool(_TRANSFER_PATTERN.search(_search_text(description, details)))


# ── Reading files ──────────────────────────────────────────────────────────────

class StatementImportError(ValueError):
    """Raised with a user-facing (Italian) message when a statement cannot be read."""


class _HTMLTableParser(HTMLParser):
    """Some banks export HTML tables with an .xls extension."""

    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def read_table(filename: str, raw: bytes) -> list[list]:
    """Return the first sheet of an .xlsx / .xls / .csv (or HTML-as-.xls) file as a list of rows."""
    if not raw:
        raise StatementImportError("Il file è vuoto.")

    if raw[:2] == b"PK":  # .xlsx is a zip archive
        import openpyxl
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise StatementImportError(f"Impossibile leggere il file Excel: {exc}")
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    if raw[:4] == b"\xd0\xcf\x11\xe0":  # legacy .xls (OLE2)
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=raw)
        except Exception as exc:
            raise StatementImportError(f"Impossibile leggere il file Excel: {exc}")
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows

    text = _decode(raw)
    if text.lstrip()[:1] == "<":
        parser = _HTMLTableParser()
        parser.feed(text)
        if parser.rows:
            return parser.rows
        raise StatementImportError("Il file non contiene tabelle leggibili.")

    if not filename.lower().endswith((".csv", ".txt")):
        raise StatementImportError("Formato non riconosciuto: carica un file .xlsx, .xls o .csv.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


# ── Layout detection ───────────────────────────────────────────────────────────

def normalize_header(value) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).split())


def _match_columns(headers: list[str], layout: BankLayout) -> dict[str, int] | None:
    """Map layout fields to column indexes; None if the header lacks the required columns."""
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for field_name in ("date", "description", "amount", "credit", "debit", "details",
                       "category", "currency", "status"):
        for alias in layout.columns.get(field_name, ()):
            index = next((i for i, h in enumerate(headers) if h == alias and i not in used), None)
            if index is not None:
                mapping[field_name] = index
                used.add(index)
                break
    has_amount = "amount" in mapping or ("credit" in mapping and "debit" in mapping)
    if "date" in mapping and "description" in mapping and has_amount:
        return mapping
    return None


@dataclass
class Layout:
    bank: BankLayout
    header_row: int
    columns: dict[str, int]
    headers: list[str] = field(default_factory=list)


def _identify_bank(rows: list[list], filename: str) -> str | None:
    preamble = _search_text(filename, *(str(c) for row in rows[:HEADER_SCAN_ROWS] for c in row if c))
    for key in DETECTION_ORDER:
        if any(marker in preamble for marker in BANKS[key].markers):
            return key
    return None


def detect_layout(rows: list[list], filename: str = "", bank: str = AUTO) -> Layout:
    if bank != AUTO and bank not in BANKS:
        raise StatementImportError(f"Banca non supportata: {bank}")

    if bank == AUTO:
        identified = _identify_bank(rows, filename)
        candidates = [identified] if identified else []
        candidates += [k for k in DETECTION_ORDER if k not in candidates]
    else:
        candidates = [bank]

    for key in candidates:
        layout = BANKS[key]
        for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
            headers = [normalize_header(c) for c in row]
            columns = _match_columns(headers, layout)
            if columns:
                return Layout(bank=layout, header_row=index, columns=columns, headers=headers)

    expected = "Data, Descrizione e Importo (oppure Entrate/Uscite)"
    raise StatementImportError(f"Intestazione dei movimenti non trovata: servono almeno le colonne {expected}.")


# ── Row parsing ────────────────────────────────────────────────────────────────

@dataclass
class StatementRow:
    date: date
    description: str
    amount: Decimal          # signed: negative = money out
    details: str | None = None
    bank_category: str | None = None
    currency: str = "EUR"
    type: str = "expense"
    category: str = "Altro"
    import_ref: str = ""
    duplicate: bool = False

    def to_dict(self) -> dict:
        return {
            "date": self.date.isoformat(),
            "description": self.description,
            "amount": str(self.amount),
            "details": self.details,
            "bank_category": self.bank_category,
            "currency": self.currency,
            "type": self.type,
            "category": self.category,
            "import_ref": self.import_ref,
            "duplicate": self.duplicate,
        }


def _cell(row: list, columns: dict[str, int], name: str):
    index = columns.get(name)
    if index is None or index >= len(row):
        return None
    value = row[index]
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _to_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return parse_date(str(value).split(" ")[0])
    except ValueError:
        return None


def _to_decimal(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    try:
        return parse_amount(str(value).replace("EUR", ""))
    except (ValueError, InvalidOperation):
        return None


def _text(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def parse_rows(rows: list[list], layout: Layout) -> tuple[list[StatementRow], int]:
    """
    Convert the rows below the header into StatementRows.
    Returns (rows, skipped) where skipped counts pending (not yet booked) movements.
    Lines without a valid date or amount (totals, balances, blank lines) are ignored.
    """
    columns = layout.columns
    parsed: list[StatementRow] = []
    skipped = 0
    for row in rows[layout.header_row + 1:]:
        tx_date = _to_date(_cell(row, columns, "date"))
        if tx_date is None:
            continue

        if "amount" in columns:
            amount = _to_decimal(_cell(row, columns, "amount"))
        else:
            credit = _to_decimal(_cell(row, columns, "credit")) or Decimal(0)
            debit = _to_decimal(_cell(row, columns, "debit")) or Decimal(0)
            amount = abs(credit) - abs(debit) if (credit or debit) else None
        if not amount:
            continue

        status = (_text(_cell(row, columns, "status")) or "").lower()
        if any(p in status for p in PENDING_STATUSES):
            skipped += 1
            continue

        description = _text(_cell(row, columns, "description"))
        details = _text(_cell(row, columns, "details"))
        if details == description:
            details = None
        if not description:
            description, details = details, None
        if not description:
            continue

        currency = (_text(_cell(row, columns, "currency")) or "EUR").upper()
        parsed.append(StatementRow(
            date=tx_date,
            description=description[:255],
            amount=amount,
            details=details,
            bank_category=_text(_cell(row, columns, "category")),
            currency=currency if re.fullmatch(r"[A-Z]{3}", currency) else "EUR",
        ))
    return parsed, skipped


def fingerprint(row: StatementRow, occurrence: int) -> str:
    """Stable id of a statement row; `occurrence` distinguishes identical rows in the same file."""
    key = f"{row.date.isoformat()}|{row.amount:.2f}|{normalize_header(row.description)}|{occurrence}"
    return hashlib.sha256(key.encode()).hexdigest()


def enrich(rows: list[StatementRow]) -> list[StatementRow]:
    """Assign type, category and import_ref, and flag rows that were already imported."""
    seen: dict[str, int] = {}
    for row in rows:
        if is_transfer(row.description, row.details):
            row.type, row.category = "transfer", "Giroconto"
        else:
            row.type = "expense" if row.amount < 0 else "income"
            row.category = categorize(row.description, row.details, row.bank_category)

        base = fingerprint(row, 0)
        occurrence = seen.get(base, 0)
        seen[base] = occurrence + 1
        row.import_ref = fingerprint(row, occurrence)

    refs = [r.import_ref for r in rows]
    existing = {
        ref for (ref,) in
        Transaction.query.with_entities(Transaction.import_ref).filter(Transaction.import_ref.in_(refs)).all()
    } if refs else set()
    for row in rows:
        row.duplicate = row.import_ref in existing
    return rows


@dataclass
class StatementPreview:
    bank: BankLayout
    rows: list[StatementRow]
    pending_skipped: int

    @property
    def new_rows(self) -> list[StatementRow]:
        return [r for r in self.rows if not r.duplicate]

    @property
    def duplicates(self) -> int:
        return sum(1 for r in self.rows if r.duplicate)

    @property
    def total_income(self) -> Decimal:
        return sum((r.amount for r in self.new_rows if r.type == "income"), Decimal(0))

    @property
    def total_expenses(self) -> Decimal:
        return sum((-r.amount for r in self.new_rows if r.type == "expense"), Decimal(0))

    @property
    def period(self) -> tuple[date, date] | None:
        if not self.rows:
            return None
        dates = [r.date for r in self.rows]
        return min(dates), max(dates)


def analyze_statement(filename: str, raw: bytes, bank: str = AUTO) -> StatementPreview:
    table = read_table(filename, raw)
    layout = detect_layout(table, filename, bank)
    rows, pending = parse_rows(table, layout)
    if not rows:
        raise StatementImportError("Nessun movimento trovato nel file.")
    rows.sort(key=lambda r: r.date)
    return StatementPreview(bank=layout.bank, rows=enrich(rows), pending_skipped=pending)


def build_transaction(data: dict, bank_key: str, category: str | None = None) -> Transaction:
    """Create a Transaction from a (serialized) StatementRow."""
    amount = Decimal(data["amount"])
    return Transaction(
        date=date.fromisoformat(data["date"]),
        description=data["description"],
        amount=abs(amount),
        currency=data.get("currency") or "EUR",
        type=data["type"],
        category=(category or data.get("category") or "Altro").strip()[:100],
        counterparty=None,
        tags=["importato", bank_key],
        is_recurring=False,
        notes=data.get("details"),
        import_ref=data["import_ref"],
    )
