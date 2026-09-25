"""
Readers that turn any supported bank-statement file into data the importer can analyze.

Supported: Excel (.xlsx, .xls, HTML saved as .xls), CSV, TXT (delimited or fixed-width), PDF (text-based),
Word (.docx), RTF, OpenDocument (.ods, .odt).

Every reader returns a `Document` with up to three views of the file, from most to least structured:
  • tables — explicit tables (spreadsheet sheets, Word/ODF tables, ruled PDF tables, delimited text)
  • lines  — words with their horizontal position, for statements laid out in columns without a real
             table (PDF text, fixed-width TXT); bank_import rebuilds the columns from the header positions
  • text_lines — plain text, for the last-resort "date … amount" line parser
"""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from statistics import median
from xml.etree import ElementTree

OLE2_MAGIC = b"\xd0\xcf\x11\xe0"
IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".heic", ".webp")


class UnsupportedFile(ValueError):
    """Raised with a user-facing (Italian) message when a file cannot be read."""


@dataclass
class Word:
    text: str
    x0: float
    x1: float
    top: float
    page: int = 0

    @property
    def center(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass
class Document:
    kind: str
    tables: list[list[list]] = field(default_factory=list)
    lines: list[list[Word]] = field(default_factory=list)
    text_lines: list[str] = field(default_factory=list)
    char_width: float = 1.0


# ── Entry point ────────────────────────────────────────────────────────────────

def read_document(filename: str, raw: bytes) -> Document:
    if not raw:
        raise UnsupportedFile("Il file è vuoto.")
    name = filename.lower()

    if raw[:5] == b"%PDF-":
        return _read_pdf(raw)
    if raw[:2] == b"PK":
        return _read_zip_document(raw)
    if raw[:4] == OLE2_MAGIC:
        if name.endswith((".doc", ".dot")):
            raise UnsupportedFile("I file Word 97-2003 (.doc) non sono supportati: aprilo in Word e salvalo come .docx o PDF.")
        return _read_xls(raw)
    if name.endswith(IMAGE_EXTENSIONS) or raw[:3] == b"\xff\xd8\xff" or raw[:8] == b"\x89PNG\r\n\x1a\n":
        raise UnsupportedFile("Le immagini non sono supportate: scarica dall'home banking il PDF, l'Excel o il CSV dei movimenti.")

    text = decode_text(raw)
    if raw[:5] == b"{\\rtf":
        return text_document("rtf", _rtf_to_text(text))
    if text.lstrip()[:1] == "<":
        return Document(kind="html", tables=[_html_rows(text)])
    return text_document("csv" if name.endswith(".csv") else "txt", text)


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


# ── Spreadsheets ───────────────────────────────────────────────────────────────

def _read_xlsx(raw: bytes) -> Document:
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedFile(f"Impossibile leggere il file Excel: {exc}")
    return Document(kind="xlsx", tables=[[list(row) for row in workbook.active.iter_rows(values_only=True)]])


def _read_xls(raw: bytes) -> Document:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:
        raise UnsupportedFile(f"Impossibile leggere il file Excel: {exc}")
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value)
        rows.append(row)
    return Document(kind="xls", tables=[rows])


def _read_zip_document(raw: bytes) -> Document:
    """.xlsx, .docx, .ods and .odt are all zip archives: tell them apart by their contents."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            names = set(archive.namelist())
            if "xl/workbook.xml" in names:
                return _read_xlsx(raw)
            if "word/document.xml" in names:
                return _read_docx(raw)
            if "content.xml" in names:
                return _read_odf(archive.read("content.xml"))
    except zipfile.BadZipFile:
        pass
    raise UnsupportedFile("Formato non riconosciuto: carica un file Excel, CSV, TXT, PDF, Word o OpenDocument.")


# ── HTML (bank "Excel" exports that are really HTML tables) ─────────────────────

def _html_rows(text: str) -> list[list[str]]:
    from html.parser import HTMLParser

    class TableParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.rows, self.row, self.cell = [], None, None

        def handle_starttag(self, tag, attrs):
            if tag == "tr":
                self.row = []
            elif tag in ("td", "th") and self.row is not None:
                self.cell = []

        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.row is not None and self.cell is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.cell = None
            elif tag == "tr" and self.row is not None:
                self.rows.append(self.row)
                self.row = None

        def handle_data(self, data):
            if self.cell is not None:
                self.cell.append(data)

    parser = TableParser()
    parser.feed(text)
    if not parser.rows:
        raise UnsupportedFile("Il file non contiene tabelle leggibili.")
    return parser.rows


# ── Plain text (CSV, TXT, RTF, Word paragraphs) ─────────────────────────────────

def text_document(kind: str, text: str) -> Document:
    """Delimited table (if any) + fixed-width positioned lines + plain lines."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    document = Document(kind=kind, text_lines=text.split("\n"), char_width=1.0)

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        rows = [row for row in csv.reader(io.StringIO(text), dialect)]
        if sum(1 for row in rows if len(row) >= 3) >= 2:
            document.tables.append(rows)
    except csv.Error:
        pass
    if "\t" in text:
        document.tables.append([line.split("\t") for line in document.text_lines])

    # Fixed-width layout: character offsets act as x coordinates, one line per row
    for number, line in enumerate(document.text_lines):
        words = [Word(m.group(), m.start(), m.end(), number) for m in re.finditer(r"\S+", line.expandtabs(8))]
        document.lines.append(words)
    return document


def _rtf_to_text(rtf: str) -> str:
    """Minimal RTF → text: keeps paragraphs, tabs and table cells, drops formatting."""
    text = re.sub(r"\{\\\*[^{}]*\}", "", rtf)                      # ignorable destinations
    text = re.sub(r"\{\\(fonttbl|colortbl|stylesheet|info)[^{}]*(\{[^{}]*\}[^{}]*)*\}", "", text)
    text = re.sub(r"\\(par|line|row)\b ?", "\n", text)
    text = re.sub(r"\\(tab|cell)\b ?", "\t", text)
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252"), text)
    text = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("\\{", "{").replace("\\}", "}").replace("\\\\", "\\")
    return re.sub(r"[{}]", "", text)


# ── Word (.docx) ───────────────────────────────────────────────────────────────

def _read_docx(raw: bytes) -> Document:
    import docx
    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise UnsupportedFile(f"Impossibile leggere il documento Word: {exc}")

    paragraphs = [p.text for p in document.paragraphs]
    preamble = [[p] for p in paragraphs[:15] if p.strip()]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells, previous = [], None
            for cell in row.cells:  # merged cells are repeated by python-docx: keep one copy
                if cell._tc is not previous:
                    cells.append(cell.text)
                previous = cell._tc
            rows.append(cells)
        tables.append(rows)

    result = text_document("docx", "\n".join(paragraphs))
    if tables:
        merged = preamble + [row for table in tables for row in table]
        result.tables = [merged] + [preamble + t for t in tables] + result.tables
    return result


# ── OpenDocument (.ods / .odt) ─────────────────────────────────────────────────

ODF = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
}


def _read_odf(content: bytes) -> Document:
    root = ElementTree.fromstring(content)
    q = lambda ns, tag: f"{{{ODF[ns]}}}{tag}"  # noqa: E731
    tables = []
    for table in root.iter(q("table", "table")):
        rows = []
        for row in table.iter(q("table", "table-row")):
            cells = []
            for cell in row:
                if cell.tag not in (q("table", "table-cell"), q("table", "covered-table-cell")):
                    continue
                value = (
                    cell.get(q("office", "date-value"))
                    or cell.get(q("office", "value"))
                    or "\n".join("".join(p.itertext()) for p in cell.iter(q("text", "p")))
                )
                repeat = min(int(cell.get(q("table", "number-columns-repeated"), "1")), 50)
                cells.extend([value] * repeat)
            while cells and not cells[-1]:
                cells.pop()
            rows.append(cells)
        tables.append(rows)
    if not tables:
        raise UnsupportedFile("Il documento non contiene tabelle con i movimenti.")
    paragraphs = [[p] for p in ("".join(e.itertext()) for e in root.iter(q("text", "p"))) if p.strip()][:10]
    merged = [row for table in tables for row in table]
    return Document(kind="odf", tables=[paragraphs + merged] + tables)


# ── PDF ────────────────────────────────────────────────────────────────────────

def _read_pdf(raw: bytes) -> Document:
    import pdfplumber
    document = Document(kind="pdf")
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            ruled_rows = []
            widths = []
            for page_number, page in enumerate(pdf.pages):
                for table in page.extract_tables():
                    ruled_rows.extend(table)
                words = page.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False)
                for w in words:
                    widths.append((w["x1"] - w["x0"]) / max(len(w["text"]), 1))
                document.lines.extend(_group_lines(
                    [Word(w["text"], w["x0"], w["x1"], w["top"], page_number) for w in words]
                ))
                document.text_lines.extend((page.extract_text() or "").split("\n"))
    except Exception as exc:
        raise UnsupportedFile(f"Impossibile leggere il PDF: {exc}")

    if not any(line.strip() for line in document.text_lines):
        raise UnsupportedFile(
            "Il PDF non contiene testo (probabilmente è una scansione): scarica dall'home banking "
            "il PDF originale oppure l'Excel/CSV dei movimenti."
        )
    document.char_width = median(widths) if widths else 4.0
    if ruled_rows:
        preamble = [[line] for line in document.text_lines[:15]]
        document.tables.append(preamble + ruled_rows)
    return document


def _group_lines(words: list[Word], tolerance: float = 2.5) -> list[list[Word]]:
    """Cluster words of one page into lines by their vertical position."""
    lines: list[list[Word]] = []
    for word in sorted(words, key=lambda w: (w.top, w.x0)):
        if lines and abs(lines[-1][0].top - word.top) <= tolerance:
            lines[-1].append(word)
        else:
            lines.append([word])
    return [sorted(line, key=lambda w: w.x0) for line in lines]
