"""
Generate FAKE bank statements for trying out the bank import (Esporta → Importa Estratto Conto Bancario).

    python samples/bank_statements/generate.py

All names, IBANs and amounts are invented. The output is deterministic (fixed random seed),
so re-running the script produces the same movements.

Requires openpyxl (in requirements.txt). The legacy binary .xls sample also needs `xlwt`
(pip install xlwt); it is skipped when xlwt is not installed.
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent
HOLDER = "MARIO ROSSI"
ORIGIN = date(2026, 1, 1)

# Merchants per category: (description shown by the bank, min €, max €)
SHOPS = {
    "grocery":   [("ESSELUNGA MILANO", 25, 140), ("CONAD CITY", 10, 60), ("LIDL ITALIA", 15, 80), ("CARREFOUR EXPRESS", 8, 45)],
    "fuel":      [("ENI STATION 1043", 40, 75), ("Q8 VIA EMILIA", 35, 70)],
    "transport": [("TRENITALIA WEB", 12, 60), ("ATM MILANO", 2.2, 2.2), ("TELEPASS SPA", 15, 40)],
    "leisure":   [("RISTORANTE DA LUIGI", 30, 90), ("PIZZERIA NAPOLI", 18, 45), ("DELIVEROO ITALY", 15, 35),
                  ("CINEMA ANTEO", 9, 20), ("BAR CENTRALE", 1.5, 6)],
    "health":    [("FARMACIA COMUNALE 3", 6, 45)],
    "other":     [("AMAZON EU SARL", 10, 120), ("DECATHLON", 20, 90), ("LIBRERIA FELTRINELLI", 9, 35)],
}


def money(rng: random.Random, low: float, high: float) -> float:
    return round(rng.uniform(low, high), 2)


def movements(start: date, end: date, seed: int, salary: float, rent: float) -> list[dict]:
    """A realistic month-by-month list of movements between start and end (inclusive).

    Each movement: date, amount (signed), short (bank operation type), full (merchant/details), kind.
    """
    rng = random.Random(seed)
    items: list[dict] = []

    def add(day: date, amount: float, short: str, full: str, kind: str):
        if start <= day <= end:
            items.append({"date": day, "amount": round(amount, 2), "short": short, "full": full, "kind": kind})

    # Always generate from the same origin so files with the same seed contain identical
    # movements where their periods overlap (like two consecutive downloads from the same account)
    month = min(ORIGIN, date(start.year, start.month, 1))
    while month <= end:
        y, m = month.year, month.month
        add(date(y, m, 1), -rent, "Bonifico SEPA", f"Bonifico a IMMOBILIARE CASA BELLA SRL per AFFITTO {m:02d}/{y}", "rent")
        add(date(y, m, 5), -17.99, "Pagamento carta", "NETFLIX.COM AMSTERDAM", "subscription")
        add(date(y, m, 6), -10.99, "Pagamento carta", "SPOTIFY AB STOCKHOLM", "subscription")
        add(date(y, m, 12), -money(rng, 55, 95), "Addebito SDD", f"ENEL ENERGIA SPA bolletta luce {m:02d}/{y}", "utility")
        add(date(y, m, 14), -money(rng, 20, 60), "Addebito SDD", "A2A ENERGIA gas", "utility")
        add(date(y, m, 18), -9.99, "Addebito SDD", "ILIAD ITALIA SPA ricarica", "subscription")
        add(date(y, m, 20), -200.00, "Giroconto", "Giroconto verso conto deposito", "transfer")
        add(date(y, m, 27), salary, "Bonifico SEPA", f"Bonifico da ACME SPA per STIPENDIO {m:02d}/{y}", "salary")
        add(date(y, m, 28) if m != 2 else date(y, m, 27), -money(rng, 1.5, 4), "Commissioni", "COMMISSIONI BONIFICO", "fee")

        # Everyday card payments
        for _ in range(rng.randint(14, 20)):
            category = rng.choices(list(SHOPS), weights=[6, 2, 2, 4, 1, 2])[0]
            name, low, high = rng.choice(SHOPS[category])
            day = date(y, m, rng.randint(1, 28))
            add(day, -money(rng, low, high), "Pagamento carta", f"PAGAMENTO POS {name}", category)

        if rng.random() < 0.5:
            add(date(y, m, rng.randint(2, 25)), money(rng, 150, 600), "Bonifico SEPA",
                "Bonifico da STUDIO BIANCHI per FATTURA consulenza", "freelance")

        month = date(y + (m == 12), m % 12 + 1, 1)

    items.sort(key=lambda i: i["date"])
    return items


def autosize(sheet):
    for column in sheet.columns:
        width = max(len(str(c.value or "")) for c in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)


# ── Fineco ─────────────────────────────────────────────────────────────────────

MONEYMAP = {
    "grocery": "Spesa", "fuel": "Carburante", "transport": "Trasporti", "leisure": "Ristoranti e bar",
    "health": "Salute", "other": "Shopping", "rent": "Affitto", "subscription": "Abbonamenti",
    "utility": "Utenze", "transfer": "Trasferimenti", "salary": "Stipendio", "fee": "Commissioni",
    "freelance": "Entrate varie",
}


def fineco(path: Path, start: date, end: date, seed: int, pending: int = 2):
    rows = movements(start, end, seed, salary=2450.00, rent=850.00)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimenti"
    ws.append(["Conto Corrente: 0012345678"])
    ws.append([f"Intestazione Conto Corrente: {HOLDER}"])
    ws.append([f"Periodo Dal: {start:%d/%m/%Y} Al: {end:%d/%m/%Y}"])
    ws.append(["Saldo Iniziale: 3.250,00"])
    ws.append([])
    ws.append(["Data_Operazione", "Data_Valuta", "Entrate", "Uscite", "Descrizione", "Descrizione_Completa", "Stato", "Moneymap"])
    for cell in ws[6]:
        cell.font = Font(bold=True)
    for index, r in enumerate(rows):
        status = "Autorizzato" if index >= len(rows) - pending else "Contabilizzato"
        day = f"{r['date']:%d/%m/%Y}"
        ws.append([
            day, day,
            r["amount"] if r["amount"] > 0 else None,
            r["amount"] if r["amount"] < 0 else None,  # Fineco shows Uscite as negative numbers
            r["short"], r["full"], status, MONEYMAP[r["kind"]],
        ])
    ws.append([])
    ws.append([f"Saldo Finale: {3250 + sum(r['amount'] for r in rows):.2f}".replace(".", ",")])
    autosize(ws)
    wb.save(path)
    return len(rows)


# ── Intesa Sanpaolo ────────────────────────────────────────────────────────────

INTESA_CATEGORIES = {
    "grocery": "Alimentari e supermercati", "fuel": "Carburanti", "transport": "Trasporti",
    "leisure": "Ristoranti e bar", "health": "Farmacie e salute", "other": "Shopping",
    "rent": "Affitto", "subscription": "Abbonamenti e servizi", "utility": "Utenze",
    "transfer": "Giroconti", "salary": "Stipendi e pensioni", "fee": "Commissioni e spese",
    "freelance": "Altre entrate",
}
INTESA_OPERATIONS = {
    "Pagamento carta": "Pagamento tramite POS", "Bonifico SEPA": "Bonifico", "Addebito SDD": "Addebito diretto",
    "Giroconto": "Giroconto", "Commissioni": "Commissioni",
}


def intesa(path: Path, start: date, end: date, seed: int):
    rows = movements(start, end, seed, salary=2180.00, rent=720.00)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista Movimenti"
    ws.append(["Lista Movimenti"])
    ws.append(["Intesa Sanpaolo - Conto corrente 1000/00098765"])
    ws.append([f"Intestatario: {HOLDER}"])
    ws.append([f"Periodo: dal {start:%d/%m/%Y} al {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["Data", "Operazione", "Dettagli", "Conto o carta", "Contabilizzazione", "Categoria ", "Valuta", "Importo"])
    for cell in ws[6]:
        cell.font = Font(bold=True)
    for index, r in enumerate(rows):
        card = r["short"] == "Pagamento carta"
        ws.append([
            r["date"],  # real Excel date cell
            INTESA_OPERATIONS[r["short"]],
            r["full"].replace("PAGAMENTO POS ", ""),
            "Carta XXXX 4321" if card else "Conto 1000/00098765",
            "Non contabilizzato" if index == len(rows) - 1 else "Contabilizzato",
            INTESA_CATEGORIES[r["kind"]],
            "EUR",
            r["amount"],
        ])
        ws.cell(row=ws.max_row, column=1).number_format = "DD/MM/YYYY"
    autosize(ws)
    wb.save(path)
    return len(rows)


def intesa_legacy_html(path: Path, start: date, end: date, seed: int):
    """Old Intesa export: an HTML table saved with the .xls extension, Accrediti/Addebiti columns."""
    rows = movements(start, end, seed, salary=2180.00, rent=720.00)

    def it(amount: float) -> str:
        return f"{abs(amount):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    lines = [
        "<html><head><meta charset='utf-8'></head><body><table border='1'>",
        "<tr><td colspan='6'>Intesa Sanpaolo - Movimenti conto corrente</td></tr>",
        f"<tr><td colspan='6'>Intestatario: {HOLDER}</td></tr>",
        "<tr><th>Data contabile</th><th>Data valuta</th><th>Descrizione</th><th>Accrediti</th><th>Addebiti</th><th>Descrizione estesa</th></tr>",
    ]
    for r in rows:
        credit = it(r["amount"]) if r["amount"] > 0 else ""
        debit = it(r["amount"]) if r["amount"] < 0 else ""
        lines.append(
            f"<tr><td>{r['date']:%d/%m/%Y}</td><td>{r['date']:%d/%m/%Y}</td><td>{INTESA_OPERATIONS[r['short']]}</td>"
            f"<td>{credit}</td><td>{debit}</td><td>{r['full']}</td></tr>"
        )
    lines.append("</table></body></html>")
    path.write_text("\n".join(lines), encoding="utf-8")
    return len(rows)


# ── Generic banks ──────────────────────────────────────────────────────────────

def unicredit_csv(path: Path, start: date, end: date, seed: int):
    """Generic Italian layout: semicolon separated, Italian number format, signed "Importo (EUR)"."""
    rows = movements(start, end, seed, salary=1980.00, rent=650.00)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Data Registrazione", "Data valuta", "Descrizione", "Importo (EUR)"])
        for r in rows:
            amount = f"{r['amount']:.2f}".replace(".", ",")
            writer.writerow([f"{r['date']:%d.%m.%Y}", f"{r['date']:%d.%m.%Y}", r["full"], amount])
    return len(rows)


def revolut_csv(path: Path, start: date, end: date, seed: int):
    """Revolut-style English CSV with signed Amount."""
    rng = random.Random(seed)
    merchants = [("Uber", 8, 25), ("Starbucks", 3, 8), ("Ryanair", 30, 150), ("Booking.com", 60, 300),
                 ("Just Eat", 12, 30), ("Amazon", 10, 80), ("Netflix", 17.99, 17.99)]
    rows, day, balance = [], start, 500.0
    while day <= end:
        if rng.random() < 0.35:
            name, low, high = rng.choice(merchants)
            amount = -money(rng, low, high)
            rows.append(("CARD_PAYMENT", day, name, amount))
        if day.day == 1:
            rows.append(("TOPUP", day, f"Payment from {HOLDER.title()}", 300.0))
        day += timedelta(days=1)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Type", "Product", "Started Date", "Completed Date", "Description", "Amount", "Fee",
                         "Currency", "State", "Balance"])
        for kind, day, description, amount in rows:
            balance += amount
            stamp = f"{day:%Y-%m-%d} 12:{rng.randint(10, 59)}:00"
            writer.writerow([kind, "Current", stamp, stamp, description, f"{amount:.2f}", "0.00", "EUR",
                             "COMPLETED", f"{balance:.2f}"])
    return len(rows)


def legacy_xls(path: Path, start: date, end: date, seed: int):
    """Binary .xls (Excel 97-2003) with a generic layout, as some smaller banks still export."""
    import xlwt
    rows = movements(start, end, seed, salary=2050.00, rent=600.00)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Movimenti")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    ws.write(0, 0, f"Estratto conto - {HOLDER}")
    for col, title in enumerate(["Data operazione", "Descrizione", "Dare", "Avere"]):
        ws.write(2, col, title)
    for index, r in enumerate(rows, start=3):
        ws.write(index, 0, r["date"], date_style)
        ws.write(index, 1, r["full"])
        if r["amount"] < 0:
            ws.write(index, 2, -r["amount"])
        else:
            ws.write(index, 3, r["amount"])
    wb.save(str(path))
    return len(rows)


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    files = [
        ("fineco_2026-06_2026-07.xlsx", fineco, (date(2026, 6, 1), date(2026, 7, 31), 11)),
        # Overlaps July with the file above → demonstrates duplicate detection
        ("fineco_2026-07_2026-09.xlsx", fineco, (date(2026, 7, 1), date(2026, 9, 24), 11)),
        ("intesa_sanpaolo_2026-04_2026-09.xlsx", intesa, (date(2026, 4, 1), date(2026, 9, 24), 22)),
        ("intesa_sanpaolo_legacy_2026-03.xls", intesa_legacy_html, (date(2026, 3, 1), date(2026, 3, 31), 33)),
        ("unicredit_2026-08_2026-09.csv", unicredit_csv, (date(2026, 8, 1), date(2026, 9, 24), 44)),
        ("revolut_2026-09.csv", revolut_csv, (date(2026, 9, 1), date(2026, 9, 24), 55)),
        ("banca_generica_2026-02.xls", legacy_xls, (date(2026, 2, 1), date(2026, 2, 28), 66)),
    ]
    for name, builder, (start, end, seed) in files:
        # The Fineco files share a seed so the overlapping July movements are identical
        try:
            count = builder(OUT / name, start, end, seed)
        except ImportError as exc:
            print(f"skip  {name}: {exc}")
            continue
        print(f"wrote {name} ({count} movimenti)")


if __name__ == "__main__":
    main()
